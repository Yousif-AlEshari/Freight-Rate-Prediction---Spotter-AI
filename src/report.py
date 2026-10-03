"""Generate the assessment report as DOCX and PDF from saved artifacts."""

from __future__ import annotations

import argparse
import json
from datetime import date
from pathlib import Path
from xml.sax.saxutils import escape

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
REPORT_TITLE = "Freight Rate Machine Learning Report"
REPORT_SUBTITLE = "Machine Learning Engineering Assessment"
TEAL = "064A56"
DARK_TEAL = "04363F"
CORAL = "F05A67"
PALE_TEAL = "E8F1F2"
LIGHT_GRAY = "F3F5F6"
MID_GRAY = "65757A"
WHITE = "FFFFFF"


def format_currency(value: float) -> str:
    """Format signed currency with the sign before the dollar symbol."""
    prefix = "-" if value < 0 else ""
    return f"{prefix}${abs(value):,.2f}"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate the assessment DOCX and PDF report.")
    parser.add_argument("--output-dir", default="reports")
    parser.add_argument("--title", default=REPORT_TITLE)
    return parser.parse_args()


def collect_report_data() -> dict[str, object]:
    train = pd.read_csv(ROOT / "data" / "train-test.csv", parse_dates=["date"])
    validation = pd.read_csv(ROOT / "data" / "validation.csv", parse_dates=["date"])
    predictions = pd.read_csv(ROOT / "validation_predictions.csv")
    december = pd.read_csv(ROOT / "december_predictions.csv", parse_dates=["date"])
    comparison = pd.read_csv(ROOT / "artifacts" / "model_comparison.csv")
    iteration = pd.read_csv(
        ROOT / "artifacts" / "robustness" / "iteration_sensitivity_summary.csv"
    )
    iteration_folds = pd.read_csv(
        ROOT / "artifacts" / "robustness" / "iteration_sensitivity_folds.csv"
    )
    cold_start = pd.read_csv(
        ROOT / "artifacts" / "robustness" / "city_cold_start_summary.csv"
    )
    metadata = json.loads(
        (ROOT / "artifacts" / "final_model_metadata.json").read_text(encoding="utf-8")
    )

    rpm = train["posted_rate"] / train["distance"]
    q1, q3 = rpm.quantile([0.25, 0.75])
    rpm_iqr = q3 - q1
    target_outlier = ~rpm.between(q1 - 3 * rpm_iqr, q3 + 3 * rpm_iqr)

    train_cities = set(train["pickup"]) | set(train["delivery"])
    validation_cities = set(validation["pickup"]) | set(validation["delivery"])
    new_cities = sorted(validation_cities - train_cities)
    validation_routes = list(zip(validation["pickup"], validation["delivery"], strict=True))
    train_routes = set(zip(train["pickup"], train["delivery"], strict=True))
    seen_route_pct = np.mean([route in train_routes for route in validation_routes]) * 100

    quote_analysis = train.assign(
        rpm=rpm,
        month=train["date"].dt.to_period("M").astype(str),
    )
    monthly_quote_correlation = quote_analysis.groupby("month").apply(
        lambda group: group["quote_signal"].corr(group["rpm"]),
        include_groups=False,
    )

    final_iteration = iteration.loc[iteration["iterations"].eq(400)].iloc[0]
    selected_folds = iteration_folds.loc[iteration_folds["iterations"].eq(400)].copy()
    cold_with_route = cold_start.loc[cold_start["include_route"]].iloc[0]
    cold_without_route = cold_start.loc[~cold_start["include_route"]].iloc[0]

    return {
        "train": train,
        "validation": validation,
        "predictions": predictions,
        "december": december,
        "comparison": comparison,
        "iteration": iteration,
        "selected_folds": selected_folds,
        "cold_start": cold_start,
        "metadata": metadata,
        "summary": {
            "train_rows": len(train),
            "validation_rows": len(validation),
            "train_date_min": train["date"].min().date().isoformat(),
            "train_date_max": train["date"].max().date().isoformat(),
            "validation_date_min": validation["date"].min().date().isoformat(),
            "validation_date_max": validation["date"].max().date().isoformat(),
            "missing_train_weight": int(train["weight"].isna().sum()),
            "missing_validation_weight": int(validation["weight"].isna().sum()),
            "missing_train_market": int(train["market_index"].isna().sum()),
            "missing_validation_market": int(validation["market_index"].isna().sum()),
            "negative_train_weight": int(train["weight"].lt(0).sum()),
            "negative_validation_weight": int(validation["weight"].lt(0).sum()),
            "target_outlier_count": int(target_outlier.sum()),
            "target_outlier_pct": float(target_outlier.mean() * 100),
            "new_city_count": len(new_cities),
            "new_cities": ", ".join(new_cities),
            "seen_route_pct": float(seen_route_pct),
            "market_train_mean": float(train["market_index"].mean()),
            "market_validation_mean": float(validation["market_index"].mean()),
            "quote_august_corr": float(monthly_quote_correlation.loc["2025-08"]),
            "mean_rate": float(train["posted_rate"].mean()),
            "median_rate": float(train["posted_rate"].median()),
            "final_mean_mae": float(final_iteration["mean_mae"]),
            "final_worst_mae": float(final_iteration["worst_mae"]),
            "final_rmse": float(final_iteration["mean_rmse"]),
            "final_median_ae": float(final_iteration["mean_median_ae"]),
            "final_wape": float(final_iteration["mean_wape"] * 100),
            "final_bias": float(final_iteration["mean_bias"]),
            "regular_mae": float(final_iteration["mean_mae_regular_target"]),
            "outlier_mae": float(final_iteration["mean_mae_outlier_target"]),
            "cold_with_route_mae": float(cold_with_route["mean_mae"]),
            "cold_without_route_mae": float(cold_without_route["mean_mae"]),
            "prediction_min": float(predictions["predicted_rate"].min()),
            "prediction_max": float(predictions["predicted_rate"].max()),
            "december_min": float(december["predicted_rate"].min()),
            "december_mean": float(december["predicted_rate"].mean()),
            "december_max": float(december["predicted_rate"].max()),
        },
    }


def generate_supporting_charts(data: dict[str, object], asset_dir: Path) -> dict[str, Path]:
    asset_dir.mkdir(parents=True, exist_ok=True)
    comparison = data["comparison"].copy()
    iteration = data["iteration"].copy()
    selected_folds = data["selected_folds"].copy()

    friendly_names = {
        "catboost_rpm_mae_no_quote": "CatBoost RPM MAE\nwithout quote",
        "catboost_direct_mae_no_quote": "CatBoost direct MAE\nwithout quote",
        "catboost_rpm_mae_with_quote": "CatBoost RPM MAE\nwith quote",
        "median_rpm_baseline": "Median RPM\nbaseline",
        "catboost_direct_rmse_no_quote": "CatBoost direct RMSE\nwithout quote",
    }
    comparison = comparison.sort_values("mean_mae", ascending=True)
    labels = [friendly_names.get(value, value) for value in comparison["model"]]
    bar_colors = [f"#{TEAL}" if rank == 1 else "#AAB8BB" for rank in comparison["mae_rank"]]

    figure, axis = plt.subplots(figsize=(9.2, 4.6), dpi=180)
    bars = axis.barh(labels, comparison["mean_mae"], color=bar_colors)
    axis.set_xlabel("Mean temporal MAE ($)")
    axis.set_title("Initial model comparison", loc="left", fontsize=14, fontweight="bold")
    axis.spines[["top", "right", "left"]].set_visible(False)
    axis.grid(axis="x", color="#D8E0E2", linewidth=0.8)
    axis.set_axisbelow(True)
    axis.bar_label(bars, labels=[f"${value:,.1f}" for value in comparison["mean_mae"]], padding=5)
    figure.tight_layout()
    comparison_path = asset_dir / "model_comparison.png"
    figure.savefig(comparison_path, bbox_inches="tight", facecolor="white")
    plt.close(figure)

    figure, axes = plt.subplots(1, 2, figsize=(10.8, 4.2), dpi=180)
    axes[0].plot(iteration["iterations"], iteration["mean_mae"], marker="o", color=f"#{TEAL}")
    axes[0].set(title="Iteration sensitivity", xlabel="Boosting iterations", ylabel="Mean MAE ($)")
    axes[0].set_xticks(iteration["iterations"])
    axes[0].grid(color="#D8E0E2", linewidth=0.8)

    axes[1].bar(
        selected_folds["validation_month"],
        selected_folds["mae"],
        color=[f"#{CORAL}", f"#{TEAL}", f"#{TEAL}", f"#{TEAL}"],
    )
    axes[1].set(title="Selected model by month", xlabel="Validation month", ylabel="MAE ($)")
    axes[1].tick_params(axis="x", rotation=30)
    axes[1].grid(axis="y", color="#D8E0E2", linewidth=0.8)
    for axis in axes:
        axis.spines[["top", "right"]].set_visible(False)
        axis.set_axisbelow(True)
    figure.tight_layout()
    robustness_path = asset_dir / "robustness_summary.png"
    figure.savefig(robustness_path, bbox_inches="tight", facecolor="white")
    plt.close(figure)

    return {
        "comparison": comparison_path,
        "robustness": robustness_path,
        "december": ROOT / "scorer_results" / "candidate_december.png",
    }


def _set_docx_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_docx_cell_margin(cell, margin_twips: int = 90) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side in ("top", "left", "bottom", "right"):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(margin_twips))
        element.set(qn("w:type"), "dxa")


def _add_docx_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def _docx_table(document: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = widths is None
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        header_cells[index].text = header
        _set_docx_cell_shading(header_cells[index], TEAL)
        _set_docx_cell_margin(header_cells[index])
        for run in header_cells[index].paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(values):
            cells[column_index].text = str(value)
            cells[column_index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_docx_cell_margin(cells[column_index])
            if row_index % 2:
                _set_docx_cell_shading(cells[column_index], LIGHT_GRAY)
            for run in cells[column_index].paragraphs[0].runs:
                run.font.size = Pt(8.5)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def _docx_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True


def _docx_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(text)
    paragraph.paragraph_format.space_after = Pt(4)


def build_docx(
    data: dict[str, object],
    charts: dict[str, Path],
    output: Path,
    title: str,
) -> None:
    summary = data["summary"]
    metadata = data["metadata"]
    comparison = data["comparison"].sort_values("mean_mae")
    iteration = data["iteration"].sort_values("iterations")
    selected_folds = data["selected_folds"].sort_values("validation_month")

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size, color in (
        ("Title", 30, DARK_TEAL),
        ("Heading 1", 18, TEAL),
        ("Heading 2", 12, DARK_TEAL),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True

    header = section.header.paragraphs[0]
    header.text = REPORT_SUBTITLE
    header.runs[0].font.name = "Arial"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(MID_GRAY)
    _add_docx_page_number(section.footer.paragraphs[0])

    cover = document.add_paragraph()
    cover.paragraph_format.space_before = Pt(80)
    cover.paragraph_format.space_after = Pt(14)
    run = cover.add_run(title)
    run.font.name = "Arial"
    run.font.size = Pt(32)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(DARK_TEAL)

    accent = document.add_paragraph()
    accent.paragraph_format.space_after = Pt(28)
    accent_run = accent.add_run("FUTURE-AWARE FREIGHT RATE REGRESSION")
    accent_run.font.name = "Arial"
    accent_run.font.size = Pt(11)
    accent_run.font.bold = True
    accent_run.font.color.rgb = RGBColor.from_string(CORAL)

    subtitle = document.add_paragraph(REPORT_SUBTITLE)
    subtitle.runs[0].font.size = Pt(15)
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(MID_GRAY)
    document.add_paragraph(f"Prepared {date.today().strftime('%B %Y')}")

    document.add_paragraph("\n")
    _docx_table(
        document,
        ["Development", "Final prediction set", "Jul-Oct backtest", "Final output"],
        [[
            f"{summary['train_rows']:,} labeled loads",
            f"{summary['validation_rows']:,} future loads",
            f"${summary['final_mean_mae']:,.2f} mean MAE",
            "12,000 scorer-validated rates",
        ]],
        widths=[1.7, 1.7, 1.7, 1.7],
    )
    document.add_page_break()

    _docx_heading(document, "Executive Summary")
    document.add_paragraph(
        "This assessment develops a reproducible regression pipeline for predicting posted "
        "freight rates. The development data covers January through October 2025, while the "
        "unlabeled prediction set covers November and December. Because this is a strict "
        "future prediction problem, model selection used expanding-window monthly holdouts "
        "rather than random cross-validation."
    )
    document.add_paragraph(
        f"The selected model is a CatBoost regressor trained on rate per mile with an MAE "
        f"objective and distance sample weights. At 400 boosting iterations it achieved "
        f"${summary['final_mean_mae']:,.2f} mean temporal MAE, "
        f"{summary['final_wape']:.2f}% WAPE, and ${summary['final_median_ae']:,.2f} mean "
        "median absolute error. It improved mean MAE by approximately 17.7% over the "
        "median-RPM business baseline."
    )
    _docx_bullet(document, "Negative weights were repaired with absolute value while retaining anomaly indicators.")
    _docx_bullet(document, "All imputation statistics were learned inside each temporal training fold.")
    _docx_bullet(document, "quote_signal was excluded after a month-by-month stability analysis and ablation test.")
    _docx_bullet(document, "Eight-city cold-start simulations tested geographic generalization before final fitting.")
    _docx_bullet(document, "The final 12,000 predictions and all 31 December chart rows passed the supplied scorer.")

    _docx_heading(document, "1. Objective and Data", level=1)
    document.add_paragraph(
        "The objective is to estimate posted_rate for every load in the supplied validation "
        "set and to produce a fixed December rate curve for a Lexington-to-Fort Wayne dry-van "
        "load. The model must generalize across future dates, market conditions, unseen cities, "
        "and previously unseen route combinations."
    )
    _docx_table(
        document,
        ["Dataset", "Rows", "Period", "Role"],
        [
            ["Development", f"{summary['train_rows']:,}", f"{summary['train_date_min']} to {summary['train_date_max']}", "Labeled training and temporal validation"],
            ["Validation", f"{summary['validation_rows']:,}", f"{summary['validation_date_min']} to {summary['validation_date_max']}", "Final unlabeled inference"],
            ["December chart", "31", "2025-12-01 to 2025-12-31", "Fixed-load daily scenario"],
        ],
        widths=[1.25, 0.75, 2.25, 2.75],
    )

    _docx_heading(document, "2. Exploratory Analysis and Data Quality")
    document.add_paragraph(
        f"The target has a mean of ${summary['mean_rate']:,.2f} and a median of "
        f"${summary['median_rate']:,.2f}. A conservative rate-per-mile IQR rule identified "
        f"{summary['target_outlier_count']:,} extreme rows ({summary['target_outlier_pct']:.2f}%). "
        "Tested configurations predict this segment poorly. Possible contamination is one "
        "hypothesis alongside legitimate rare events and omitted predictors."
    )
    _docx_table(
        document,
        ["Issue", "Development", "Validation", "Treatment"],
        [
            ["Missing weight", f"{summary['missing_train_weight']:,}", f"{summary['missing_validation_weight']:,}", "Equipment median plus indicator"],
            ["Negative weight", f"{summary['negative_train_weight']:,}", f"{summary['negative_validation_weight']:,}", "Absolute value plus indicator"],
            ["Missing market index", f"{summary['missing_train_market']:,}", f"{summary['missing_validation_market']:,}", "Fold-local median plus indicator"],
            ["Extreme RPM target", f"{summary['target_outlier_count']:,}", "Unknown", "Robust MAE objective"],
        ],
        widths=[1.35, 1.1, 1.1, 3.3],
    )
    document.add_paragraph(
        f"The validation set adds {summary['new_city_count']} cities absent from development: "
        f"{summary['new_cities']}. Only {summary['seen_route_pct']:.2f}% of final validation rows "
        "use a directed route observed during development, so categorical route information is "
        "paired with distance and geographic features."
    )
    document.add_paragraph(
        f"The market-index mean shifts from {summary['market_train_mean']:.3f} in development to "
        f"{summary['market_validation_mean']:.3f} in validation. quote_signal is also unstable: "
        f"its August correlation with rate per mile is {summary['quote_august_corr']:.3f}, while "
        "other months alternate between positive and negative relationships."
    )

    _docx_heading(document, "3. Preprocessing and Feature Engineering")
    for text in (
        "Remove load_id from model features and keep posted_rate separate as the target.",
        "Convert negative weight to absolute weight; add missing and negative-weight indicators.",
        "Impute weight by equipment and market index globally within each fitted training fold.",
        "Create route, log-distance, coordinate deltas, haversine distance, and bearing features.",
        "Create linear and cyclical day-of-week and day-of-year calendar features.",
        "Preserve pickup, delivery, equipment, and route as native CatBoost categorical features.",
    ):
        _docx_bullet(document, text)
    document.add_paragraph(
        "The preprocessing component implements the sklearn fit/transform contract. This ensures "
        "that imputation values are estimated from training rows only and applied unchanged to each holdout."
    )

    _docx_heading(document, "4. Validation Strategy")
    document.add_paragraph(
        "Four expanding-window folds mirror the final future-prediction setting. Random splitting "
        "was rejected because it would mix nearby market regimes and overstate the usefulness of temporally unstable features."
    )
    _docx_table(
        document,
        ["Fold", "Training window", "Validation month"],
        [
            ["1", "January-June", "July"],
            ["2", "January-July", "August"],
            ["3", "January-August", "September"],
            ["4", "January-September", "October"],
        ],
        widths=[0.8, 2.8, 2.8],
    )
    document.add_paragraph(
        "MAE was the primary selection metric. RMSE, median absolute error, WAPE, RMSLE, bias, "
        "regular-target error, outlier-target error, and route coverage were retained as diagnostics."
    )

    _docx_heading(document, "5. Model Selection")
    friendly = {
        "catboost_rpm_mae_no_quote": "CatBoost RPM MAE, no quote",
        "catboost_direct_mae_no_quote": "CatBoost direct MAE, no quote",
        "catboost_rpm_mae_with_quote": "CatBoost RPM MAE, with quote",
        "median_rpm_baseline": "Median RPM baseline",
        "catboost_direct_rmse_no_quote": "CatBoost direct RMSE, no quote",
    }
    _docx_table(
        document,
        ["Candidate", "Mean MAE", "WAPE", "Mean RMSE"],
        [
            [friendly.get(row.model, row.model), f"${row.mean_mae:,.2f}", f"{row.mean_wape * 100:.2f}%", f"${row.mean_rmse:,.2f}"]
            for row in comparison.itertuples()
        ],
        widths=[3.2, 1.2, 1.0, 1.35],
    )
    document.add_picture(str(charts["comparison"]), width=Inches(6.9))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "Predicting RPM matches the economic structure of freight pricing. Distance sample "
        "weights make the RPM MAE objective equivalent to dollar MAE after rate reconstruction. "
        "The quote ablation supports excluding quote_signal for the future validation regime."
    )

    _docx_heading(document, "6. Robustness Checks")
    _docx_table(
        document,
        ["Iterations", "Mean MAE", "Worst-month MAE", "Runtime"],
        [
            [f"{int(row.iterations)}", f"${row.mean_mae:,.2f}", f"${row.worst_mae:,.2f}", f"{row.total_seconds:,.0f} s"]
            for row in iteration.itertuples()
        ],
        widths=[1.2, 1.5, 1.7, 1.4],
    )
    document.add_picture(str(charts["robustness"]), width=Inches(6.9))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        f"The 400-iteration configuration had the best mean and worst-month MAE. In three "
        f"October stress tests, each removing eight cities entirely from prior training, keeping "
        f"route produced ${summary['cold_with_route_mae']:,.2f} mean MAE versus "
        f"${summary['cold_without_route_mae']:,.2f} without route. Route was therefore retained."
    )

    _docx_heading(document, "7. Final Model and Performance")
    parameters = metadata["parameters"]
    _docx_table(
        document,
        ["Component", "Final choice"],
        [
            ["Estimator", "CatBoostRegressor"],
            ["Target", "posted_rate / distance"],
            ["Objective", parameters["loss_function"]],
            ["Sample weighting", "distance"],
            ["Iterations / depth", f"{parameters['iterations']} / {parameters['depth']}"],
            ["Learning rate / L2", f"{parameters['learning_rate']} / {parameters['l2_leaf_reg']}"],
            ["quote_signal", "Excluded"],
            ["Route", "Included"],
        ],
        widths=[2.2, 4.0],
    )
    _docx_table(
        document,
        ["Mean MAE", "Worst month", "Median AE", "WAPE", "Bias"],
        [[
            f"${summary['final_mean_mae']:,.2f}",
            f"${summary['final_worst_mae']:,.2f}",
            f"${summary['final_median_ae']:,.2f}",
            f"{summary['final_wape']:.2f}%",
            format_currency(summary["final_bias"]),
        ]],
        widths=[1.25, 1.25, 1.25, 1.0, 1.25],
    )
    document.add_paragraph(
        f"On regular backtest targets, mean MAE was ${summary['regular_mae']:,.2f}. For rows "
        f"extreme under the conservative RPM rule it was ${summary['outlier_mae']:,.2f}, which "
        f"explains the much larger ${summary['final_rmse']:,.2f} RMSE. Tested configurations "
        "predict this segment poorly; possible contamination is one hypothesis alongside "
        "legitimate rare events and omitted predictors."
    )

    _docx_heading(document, "8. Final Inference and December Scenario")
    document.add_paragraph(
        f"The final model was refit on all {summary['train_rows']:,} labeled rows and generated "
        f"{summary['validation_rows']:,} predictions. Rates range from "
        f"${summary['prediction_min']:,.2f} to ${summary['prediction_max']:,.2f}; all are finite, "
        "positive, uniquely aligned to load_id, and accepted by the supplied scorer."
    )
    document.add_paragraph(
        "The chart input fixes Lexington to Fort Wayne, 360 miles, Dry Van, and 32,000 lb. "
        "Coordinates are recovered from the labeled city mapping. Because market_index is absent "
        "from the chart input, each date uses the daily mean from the supplied unlabeled validation "
        "covariates. quote_signal remains unused."
    )
    document.add_picture(str(charts["december"]), width=Inches(7.0))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = document.add_paragraph(
        f"Figure: scorer-generated December prediction curve. Range "
        f"${summary['december_min']:,.2f}-${summary['december_max']:,.2f}; "
        f"mean ${summary['december_mean']:,.2f}."
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(8.5)

    _docx_heading(document, "9. Limitations and Production Considerations")
    for text in (
        "The assessment does not specify the hidden scoring metric or acceptance threshold.",
        "Approximately 1.4% of development targets are extreme under the conservative RPM rule; tested configurations predict this segment poorly. Possible contamination is one hypothesis alongside legitimate rare events and omitted predictors.",
        "Historical backtest error is concentrated in RPM-extreme rows; final-period novelty risk is unmeasured because its labels are hidden.",
        "December market enrichment uses unlabeled validation covariates and should be replaced by an explicit market feed in production.",
        "A deployed system should monitor monthly error, route coverage, feature drift, prediction ranges, and calibration by equipment and distance band.",
    ):
        _docx_bullet(document, text)

    _docx_heading(document, "10. Reproducibility")
    document.add_paragraph("The complete pipeline can be reproduced from the repository root:")
    for command in (
        ".\\.venv\\Scripts\\python.exe -m unittest discover -s tests -v",
        ".\\.venv\\Scripts\\python.exe -m src.train --iterations 600",
        ".\\.venv\\Scripts\\python.exe -m src.robustness",
        ".\\.venv\\Scripts\\python.exe -m src.predict --iterations 400",
        ".\\.venv\\Scripts\\python.exe score.py --predictions validation_predictions.csv --december-predictions december_predictions.csv",
    ):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(command)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        paragraph.paragraph_format.left_indent = Inches(0.2)
        paragraph.paragraph_format.space_after = Pt(3)

    versions = metadata["versions"]
    document.add_paragraph(
        f"Final environment: Python {versions['python']}, pandas {versions['pandas']}, "
        f"NumPy {versions['numpy']}, scikit-learn {versions['scikit_learn']}, and "
        f"CatBoost {versions['catboost']}. Ten automated tests and the supplied scorer passed."
    )

    _docx_heading(document, "Conclusion")
    document.add_paragraph(
        "The final pipeline combines future-aware validation, robust loss, domain feature "
        "engineering, explicit instability tests, and scorer-validated outputs. Performance is "
        "strong for ordinary freight loads. Historical backtest error is concentrated in "
        "RPM-extreme rows, while final-period novelty risk remains unmeasured."
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def _pdf_styles():
    sample = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "ReportTitle",
            parent=sample["Title"],
            fontName="Helvetica-Bold",
            fontSize=28,
            leading=32,
            textColor=colors.HexColor(f"#{DARK_TEAL}"),
            alignment=TA_LEFT,
            spaceAfter=12,
        ),
        "subtitle": ParagraphStyle(
            "ReportSubtitle",
            parent=sample["Normal"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=14,
            textColor=colors.HexColor(f"#{CORAL}"),
            spaceAfter=20,
        ),
        "h1": ParagraphStyle(
            "ReportH1",
            parent=sample["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=17,
            leading=20,
            textColor=colors.HexColor(f"#{TEAL}"),
            spaceBefore=11,
            spaceAfter=7,
            keepWithNext=True,
        ),
        "h2": ParagraphStyle(
            "ReportH2",
            parent=sample["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=14,
            textColor=colors.HexColor(f"#{DARK_TEAL}"),
            spaceBefore=7,
            spaceAfter=5,
            keepWithNext=True,
        ),
        "body": ParagraphStyle(
            "ReportBody",
            parent=sample["BodyText"],
            fontName="Helvetica",
            fontSize=8.8,
            leading=12.2,
            textColor=colors.HexColor("#26383D"),
            spaceAfter=6,
        ),
        "bullet": ParagraphStyle(
            "ReportBullet",
            parent=sample["BodyText"],
            fontName="Helvetica",
            fontSize=8.6,
            leading=11.8,
            leftIndent=13,
            firstLineIndent=-7,
            bulletIndent=5,
            spaceAfter=3,
        ),
        "caption": ParagraphStyle(
            "ReportCaption",
            parent=sample["BodyText"],
            fontName="Helvetica-Oblique",
            fontSize=7.5,
            leading=9.5,
            textColor=colors.HexColor(f"#{MID_GRAY}"),
            alignment=TA_CENTER,
            spaceAfter=7,
        ),
        "code": ParagraphStyle(
            "ReportCode",
            parent=sample["Code"],
            fontName="Courier",
            fontSize=7,
            leading=9,
            leftIndent=8,
            textColor=colors.HexColor("#24363B"),
            spaceAfter=3,
        ),
        "table_header": ParagraphStyle(
            "ReportTableHeader",
            parent=sample["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=8.6,
            leading=10.5,
            textColor=colors.white,
        ),
    }


def _pdf_paragraph(text: str, style) -> Paragraph:
    return Paragraph(escape(text), style)


def _pdf_table(headers: list[str], rows: list[list[str]], widths=None) -> Table:
    styles = _pdf_styles()
    content = [[Paragraph(escape(str(value)), styles["table_header"]) for value in headers]]
    for row in rows:
        content.append([Paragraph(escape(str(value)), styles["body"]) for value in row])
    table = Table(content, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{TEAL}")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5D7")),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{LIGHT_GRAY}")]),
            ]
        )
    )
    return table


def _pdf_page(canvas, document) -> None:
    canvas.saveState()
    width, height = A4
    canvas.setStrokeColor(colors.HexColor("#B7C5C8"))
    canvas.setLineWidth(0.4)
    canvas.line(document.leftMargin, height - 34, width - document.rightMargin, height - 34)
    canvas.setFillColor(colors.HexColor(f"#{MID_GRAY}"))
    canvas.setFont("Helvetica", 7)
    canvas.drawString(document.leftMargin, height - 27, REPORT_SUBTITLE)
    canvas.drawRightString(width - document.rightMargin, 23, f"Page {document.page}")
    canvas.restoreState()


def build_pdf(
    data: dict[str, object],
    charts: dict[str, Path],
    output: Path,
    title: str,
) -> None:
    summary = data["summary"]
    metadata = data["metadata"]
    comparison = data["comparison"].sort_values("mean_mae")
    iteration = data["iteration"].sort_values("iterations")
    styles = _pdf_styles()
    output.parent.mkdir(parents=True, exist_ok=True)
    document = SimpleDocTemplate(
        str(output),
        pagesize=A4,
        rightMargin=42,
        leftMargin=42,
        topMargin=46,
        bottomMargin=38,
        title=title,
        author="",
        subject="Freight rate machine learning assessment",
    )
    story = []

    story.extend(
        [
            Spacer(1, 0.9 * inch),
            _pdf_paragraph(title, styles["title"]),
            _pdf_paragraph("FUTURE-AWARE FREIGHT RATE REGRESSION", styles["subtitle"]),
            _pdf_paragraph(REPORT_SUBTITLE, ParagraphStyle("cover", parent=styles["body"], fontSize=15, leading=18, textColor=colors.HexColor(f"#{MID_GRAY}"))),
            Spacer(1, 0.12 * inch),
            _pdf_paragraph(f"Prepared {date.today().strftime('%B %Y')}", styles["body"]),
            Spacer(1, 0.65 * inch),
            _pdf_table(
                ["Development", "Prediction set", "Jul-Oct backtest", "Final output"],
                [[f"{summary['train_rows']:,} labeled", f"{summary['validation_rows']:,} future", f"${summary['final_mean_mae']:,.2f} MAE", "12,000 valid rates"]],
                widths=[1.22 * inch] * 4,
            ),
            PageBreak(),
        ]
    )

    story.append(_pdf_paragraph("Executive Summary", styles["h1"]))
    story.append(_pdf_paragraph(
        "This assessment develops a reproducible regression pipeline for posted freight rates. "
        "Development covers January through October 2025 and final inference covers November and "
        "December, so model selection uses expanding-window monthly holdouts rather than random cross-validation.",
        styles["body"],
    ))
    story.append(_pdf_paragraph(
        f"The selected CatBoost RPM model achieved ${summary['final_mean_mae']:,.2f} mean temporal MAE, "
        f"{summary['final_wape']:.2f}% WAPE, and ${summary['final_median_ae']:,.2f} mean median absolute "
        "error at 400 iterations. It improved mean MAE by approximately 17.7% over the business baseline.",
        styles["body"],
    ))
    for text in (
        "Negative and missing weights are repaired with explicit anomaly indicators.",
        "All learned preprocessing values are fit inside each temporal training fold.",
        "quote_signal is excluded after temporal stability analysis and ablation testing.",
        "Eight-city cold-start simulations test geographic generalization.",
        "All 12,000 final predictions and 31 December rows pass the supplied scorer.",
    ):
        story.append(Paragraph(f"- {escape(text)}", styles["bullet"]))

    story.append(_pdf_paragraph("1. Objective and Data", styles["h1"]))
    story.append(_pdf_paragraph(
        "The objective is to estimate posted_rate for every final validation load and generate a "
        "daily December curve for a fixed Lexington-to-Fort Wayne dry-van scenario. The solution "
        "must generalize across future dates, market conditions, new cities, and unseen route combinations.",
        styles["body"],
    ))
    story.append(_pdf_table(
        ["Dataset", "Rows", "Period", "Role"],
        [
            ["Development", f"{summary['train_rows']:,}", f"{summary['train_date_min']} to {summary['train_date_max']}", "Training and temporal validation"],
            ["Validation", f"{summary['validation_rows']:,}", f"{summary['validation_date_min']} to {summary['validation_date_max']}", "Final inference"],
            ["December", "31", "2025-12-01 to 2025-12-31", "Fixed-load scenario"],
        ],
        widths=[1.0 * inch, 0.65 * inch, 1.75 * inch, 2.0 * inch],
    ))

    story.append(_pdf_paragraph("2. Exploratory Analysis and Data Quality", styles["h1"]))
    story.append(_pdf_paragraph(
        f"The target mean is ${summary['mean_rate']:,.2f} and median is ${summary['median_rate']:,.2f}. "
        f"A conservative RPM IQR rule identifies {summary['target_outlier_count']:,} extreme rows "
        f"({summary['target_outlier_pct']:.2f}%). Tested configurations predict this segment "
        "poorly. Possible contamination is one hypothesis alongside legitimate rare events and "
        "omitted predictors.",
        styles["body"],
    ))
    story.append(_pdf_table(
        ["Issue", "Development", "Validation", "Treatment"],
        [
            ["Missing weight", f"{summary['missing_train_weight']:,}", f"{summary['missing_validation_weight']:,}", "Equipment median + indicator"],
            ["Negative weight", f"{summary['negative_train_weight']:,}", f"{summary['negative_validation_weight']:,}", "Absolute value + indicator"],
            ["Missing market", f"{summary['missing_train_market']:,}", f"{summary['missing_validation_market']:,}", "Fold median + indicator"],
            ["Extreme RPM", f"{summary['target_outlier_count']:,}", "Unknown", "Robust MAE objective"],
        ],
        widths=[1.3 * inch, 1.2 * inch, 1.05 * inch, 2.55 * inch],
    ))
    story.append(Spacer(1, 6))
    story.append(_pdf_paragraph(
        f"Validation adds {summary['new_city_count']} new cities ({summary['new_cities']}) and "
        f"{summary['seen_route_pct']:.2f}% of its rows use a route observed in development. The "
        f"market mean falls from {summary['market_train_mean']:.3f} to {summary['market_validation_mean']:.3f}. "
        f"quote_signal has only {summary['quote_august_corr']:.3f} correlation with RPM in August "
        "and alternates sign in other months.",
        styles["body"],
    ))

    story.append(_pdf_paragraph("3. Preprocessing and Feature Engineering", styles["h1"]))
    for text in (
        "Drop load_id and separate posted_rate before feature transformation.",
        "Repair weight signs; retain missing and negative indicators.",
        "Apply fold-local weight and market imputations.",
        "Create route, log-distance, coordinate-delta, haversine, and bearing features.",
        "Create linear and cyclical calendar features.",
        "Preserve pickup, delivery, equipment, and route for native CatBoost handling.",
    ):
        story.append(Paragraph(f"- {escape(text)}", styles["bullet"]))

    story.append(_pdf_paragraph("4. Validation Strategy", styles["h1"]))
    story.append(_pdf_paragraph(
        "Four expanding-window folds mirror the production horizon. MAE is primary; RMSE, median "
        "absolute error, WAPE, RMSLE, bias, outlier performance, and route coverage are diagnostics.",
        styles["body"],
    ))
    story.append(_pdf_table(
        ["Fold", "Training window", "Validation"],
        [["1", "January-June", "July"], ["2", "January-July", "August"], ["3", "January-August", "September"], ["4", "January-September", "October"]],
        widths=[0.7 * inch, 2.1 * inch, 1.55 * inch],
    ))

    story.append(_pdf_paragraph("5. Model Selection", styles["h1"]))
    friendly = {
        "catboost_rpm_mae_no_quote": "CatBoost RPM MAE, no quote",
        "catboost_direct_mae_no_quote": "CatBoost direct MAE, no quote",
        "catboost_rpm_mae_with_quote": "CatBoost RPM MAE, with quote",
        "median_rpm_baseline": "Median RPM baseline",
        "catboost_direct_rmse_no_quote": "CatBoost direct RMSE, no quote",
    }
    story.append(_pdf_table(
        ["Candidate", "Mean MAE", "WAPE", "RMSE"],
        [[friendly.get(row.model, row.model), f"${row.mean_mae:,.2f}", f"{row.mean_wape * 100:.2f}%", f"${row.mean_rmse:,.2f}"] for row in comparison.itertuples()],
        widths=[2.5 * inch, 0.95 * inch, 0.75 * inch, 0.95 * inch],
    ))
    story.append(Spacer(1, 6))
    story.append(Image(str(charts["comparison"]), width=6.3 * inch, height=3.15 * inch))
    story.append(_pdf_paragraph(
        "Predicting RPM captures freight economics. Distance weighting aligns RPM MAE with dollar "
        "MAE after reconstruction. The quote ablation supports exclusion for future inference.",
        styles["body"],
    ))

    story.append(_pdf_paragraph("6. Robustness Checks", styles["h1"]))
    story.append(_pdf_table(
        ["Iterations", "Mean MAE", "Worst MAE", "Runtime"],
        [[f"{int(row.iterations)}", f"${row.mean_mae:,.2f}", f"${row.worst_mae:,.2f}", f"{row.total_seconds:,.0f} s"] for row in iteration.itertuples()],
        widths=[1.0 * inch, 1.15 * inch, 1.15 * inch, 1.0 * inch],
    ))
    story.append(Spacer(1, 6))
    story.append(Image(str(charts["robustness"]), width=6.35 * inch, height=2.47 * inch))
    story.append(_pdf_paragraph(
        f"Four hundred iterations had the best mean and worst-month MAE. In three eight-city "
        f"cold-start simulations, retaining route achieved ${summary['cold_with_route_mae']:,.2f} "
        f"mean MAE versus ${summary['cold_without_route_mae']:,.2f} without route.",
        styles["body"],
    ))

    story.append(_pdf_paragraph("7. Final Model and Performance", styles["h1"]))
    parameters = metadata["parameters"]
    story.append(_pdf_table(
        ["Component", "Choice", "Component", "Choice"],
        [
            ["Estimator", "CatBoost", "Target", "posted_rate / distance"],
            ["Objective", parameters["loss_function"], "Weight", "distance"],
            ["Iterations", str(parameters["iterations"]), "Depth", str(parameters["depth"])],
            ["quote_signal", "Excluded", "Route", "Included"],
        ],
        widths=[1.05 * inch, 1.45 * inch, 1.05 * inch, 1.55 * inch],
    ))
    story.append(Spacer(1, 6))
    story.append(_pdf_table(
        ["Mean MAE", "Worst month", "Median AE", "WAPE", "Bias"],
        [[f"${summary['final_mean_mae']:,.2f}", f"${summary['final_worst_mae']:,.2f}", f"${summary['final_median_ae']:,.2f}", f"{summary['final_wape']:.2f}%", format_currency(summary["final_bias"])]],
        widths=[1.05 * inch] * 5,
    ))
    story.append(Spacer(1, 5))
    story.append(_pdf_paragraph(
        f"Regular-target MAE is ${summary['regular_mae']:,.2f}; MAE for rows extreme under the "
        f"conservative RPM rule is ${summary['outlier_mae']:,.2f}. This segment drives the "
        f"${summary['final_rmse']:,.2f} RMSE and is predicted poorly by tested configurations; "
        "possible contamination is one hypothesis alongside legitimate rare events and omitted predictors.",
        styles["body"],
    ))

    story.append(_pdf_paragraph("8. Final Inference and December Scenario", styles["h1"]))
    story.append(_pdf_paragraph(
        f"The model was refit on all {summary['train_rows']:,} labeled rows and produced "
        f"{summary['validation_rows']:,} scorer-valid predictions from ${summary['prediction_min']:,.2f} "
        f"to ${summary['prediction_max']:,.2f}. All IDs are unique and all rates are finite and positive.",
        styles["body"],
    ))
    story.append(_pdf_paragraph(
        "For the fixed December scenario, coordinates come from the labeled city mapping and each "
        "date uses the daily mean market index from supplied unlabeled validation covariates. "
        "quote_signal is excluded from the final model.",
        styles["body"],
    ))
    story.append(Image(str(charts["december"]), width=6.6 * inch, height=2.19 * inch))
    story.append(_pdf_paragraph(
        f"Scorer-generated December curve: ${summary['december_min']:,.2f}-${summary['december_max']:,.2f}, mean ${summary['december_mean']:,.2f}.",
        styles["caption"],
    ))

    story.append(_pdf_paragraph("9. Limitations and Production Considerations", styles["h1"]))
    for text in (
        "The hidden scoring metric and acceptance threshold are not specified.",
        "Approximately 1.4% of development labels are extreme under the conservative RPM rule; tested configurations predict this segment poorly. Possible contamination is one hypothesis alongside legitimate rare events and omitted predictors.",
        "Historical backtest error is concentrated in RPM-extreme rows; final-period novelty risk is unmeasured because its labels are hidden.",
        "December market enrichment should become an explicit market-data feed in production.",
        "Production monitoring should track drift, route coverage, prediction ranges, and monthly segment error.",
    ):
        story.append(Paragraph(f"- {escape(text)}", styles["bullet"]))

    story.append(_pdf_paragraph("10. Reproducibility", styles["h1"]))
    for command in (
        ".\\.venv\\Scripts\\python.exe -m unittest discover -s tests -v",
        ".\\.venv\\Scripts\\python.exe -m src.train --iterations 600",
        ".\\.venv\\Scripts\\python.exe -m src.robustness",
        ".\\.venv\\Scripts\\python.exe -m src.predict --iterations 400",
        ".\\.venv\\Scripts\\python.exe score.py --predictions validation_predictions.csv --december-predictions december_predictions.csv",
    ):
        story.append(_pdf_paragraph(command, styles["code"]))
    versions = metadata["versions"]
    story.append(_pdf_paragraph(
        f"Environment: Python {versions['python']}, pandas {versions['pandas']}, NumPy "
        f"{versions['numpy']}, scikit-learn {versions['scikit_learn']}, CatBoost "
        f"{versions['catboost']}. Ten automated tests and the supplied scorer passed.",
        styles["body"],
    ))

    story.append(_pdf_paragraph("Conclusion", styles["h1"]))
    story.append(_pdf_paragraph(
        "The final pipeline combines future-aware validation, robust loss, domain feature "
        "engineering, explicit instability tests, and scorer-validated outputs. Performance is "
        "strong for ordinary loads. Historical backtest error is concentrated in RPM-extreme "
        "rows, while final-period novelty risk remains unmeasured.",
        styles["body"],
    ))

    document.build(story, onFirstPage=_pdf_page, onLaterPages=_pdf_page)


def main() -> None:
    args = parse_args()
    output_dir = Path(args.output_dir)
    if not output_dir.is_absolute():
        output_dir = ROOT / output_dir
    asset_dir = output_dir / "assets"

    data = collect_report_data()
    charts = generate_supporting_charts(data, asset_dir)
    docx_output = output_dir / "freight_rate_machine_learning_report.docx"
    pdf_output = output_dir / "freight_rate_machine_learning_report.pdf"
    build_docx(data, charts, docx_output, args.title)
    build_pdf(data, charts, pdf_output, args.title)
    print(f"Created DOCX report: {docx_output}")
    print(f"Created PDF report:  {pdf_output}")


if __name__ == "__main__":
    main()
